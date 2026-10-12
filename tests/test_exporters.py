import json
from pathlib import Path

from openpyxl import load_workbook

from invoice_extractor.config import Settings
from invoice_extractor.pipeline import export_all, process_files


def test_excel_and_json_export(samples_dir: Path, settings: Settings) -> None:
    invoices, failures = process_files([samples_dir], settings)
    assert not failures
    excel, js = export_all(invoices, settings)

    payload = json.loads(js.read_text(encoding="utf-8"))
    assert payload["count"] == 10
    assert payload["valid"] == 9
    first = payload["invoices"][0]
    assert first["source_file"] == "invoice_01_northwind.pdf"
    assert first["validation"]["is_valid"] is True
    assert isinstance(first["line_items"][0]["amount"], float)

    wb = load_workbook(excel)
    assert wb.sheetnames == ["Invoices", "Line Items"]
    summary = wb["Invoices"]
    assert summary.max_row == 11  # header + 10 invoices
    assert summary["A1"].font.bold
    assert summary.freeze_panes == "A2"
    assert summary["L2"].number_format == "#,##0.00"
    valid_column = [row[12].value for row in summary.iter_rows(min_row=2)]
    assert valid_column.count("NO") == 1
    bad_row = next(r for r in summary.iter_rows(min_row=2) if r[12].value == "NO")
    assert bad_row[0].value == "invoice_10_copperleaf.pdf"
    assert bad_row[0].fill.fgColor.rgb.endswith("F8D7DA")

    items = wb["Line Items"]
    assert items.max_row == sum(len(inv.line_items) for inv in invoices) + 1


def test_export_json_only(samples_dir: Path, settings: Settings) -> None:
    invoices, _ = process_files([samples_dir / "invoice_05_alderfinch.pdf"], settings)
    written = export_all(invoices, settings, formats=("json",))
    assert [p.suffix for p in written] == [".json"]
    data = json.loads(written[0].read_text(encoding="utf-8"))
    assert data["invoices"][0]["currency"] == "GBP"
