"""Export extracted invoices to JSON and a formatted Excel workbook."""

from __future__ import annotations

import json
import logging
from collections.abc import Iterable, Sequence
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import Invoice

log = logging.getLogger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
INVALID_FILL = PatternFill("solid", fgColor="F8D7DA")
VALID_FONT = Font(color="1E7B34", bold=True)
INVALID_FONT = Font(color="A61B29", bold=True)
THIN = Side(style="thin", color="D0D7DE")
BORDER = Border(bottom=THIN)
MONEY_FORMAT = "#,##0.00"
DATE_FORMAT = "yyyy-mm-dd"


def export_json(invoices: Sequence[Invoice], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "count": len(invoices),
        "valid": sum(inv.validation.is_valid for inv in invoices),
        "invoices": [inv.to_dict() for inv in invoices],
    }
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    log.info("wrote %s", path)
    return path


def _write_header(ws: Worksheet, headers: Iterable[str]) -> None:
    ws.append(list(headers))
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


def _autosize(ws: Worksheet, max_width: int = 60) -> None:
    for column_cells in ws.columns:
        length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(length + 2, 10), max_width)


def _format_columns(ws: Worksheet, money_cols: Iterable[int], date_cols: Iterable[int] = ()) -> None:
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.border = BORDER
        for idx in money_cols:
            row[idx - 1].number_format = MONEY_FORMAT
        for idx in date_cols:
            row[idx - 1].number_format = DATE_FORMAT


SUMMARY_HEADERS = (
    "Source File", "Invoice #", "Invoice Date", "Due Date", "Vendor", "Currency", "Items",
    "Subtotal", "Tax Rate %", "Tax", "Shipping", "Total", "Valid", "Issues",
)
ITEM_HEADERS = ("Source File", "Invoice #", "Vendor", "Line", "Description", "Qty", "Unit Price", "Amount", "Currency")


def export_excel(invoices: Sequence[Invoice], path: Path) -> Path:
    """Write a two-sheet workbook: an invoice summary and all line items."""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    summary = wb.active
    summary.title = "Invoices"
    _write_header(summary, SUMMARY_HEADERS)
    for inv in invoices:
        issues = "; ".join(inv.validation.errors + [f"warning: {w}" for w in inv.validation.warnings])
        summary.append([
            inv.source_file, inv.invoice_number, inv.invoice_date, inv.due_date, inv.vendor, inv.currency,
            len(inv.line_items),
            float(inv.subtotal) if inv.subtotal is not None else None,
            float(inv.tax_rate) if inv.tax_rate is not None else None,
            float(inv.tax) if inv.tax is not None else None,
            float(inv.shipping),
            float(inv.total) if inv.total is not None else None,
            "YES" if inv.validation.is_valid else "NO",
            issues,
        ])
        row = summary[summary.max_row]
        valid_cell = row[SUMMARY_HEADERS.index("Valid")]
        valid_cell.font = VALID_FONT if inv.validation.is_valid else INVALID_FONT
        valid_cell.alignment = Alignment(horizontal="center")
        if not inv.validation.is_valid:
            for cell in row:
                cell.fill = INVALID_FILL
    _format_columns(summary, money_cols=(8, 10, 11, 12), date_cols=(3, 4))
    summary.auto_filter.ref = summary.dimensions
    _autosize(summary)

    items = wb.create_sheet("Line Items")
    _write_header(items, ITEM_HEADERS)
    for inv in invoices:
        for n, item in enumerate(inv.line_items, start=1):
            items.append([
                inv.source_file, inv.invoice_number, inv.vendor, n, item.description,
                float(item.quantity), float(item.unit_price), float(item.amount), inv.currency,
            ])
    _format_columns(items, money_cols=(7, 8))
    items.auto_filter.ref = items.dimensions
    _autosize(items)

    wb.properties.title = "Extracted invoices"
    wb.properties.creator = "invoice-data-extractor"
    wb.save(path)
    log.info("wrote %s", path)
    return path
